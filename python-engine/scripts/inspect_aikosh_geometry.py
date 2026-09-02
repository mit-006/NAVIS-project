"""
Inspect SHRUG/AIKOSH PC11 Village Polygon Geometry for Kamrup Metropolitan.

This script:
1. Reads the downloaded village polygon shapefile or GeoPackage
2. Filters to Kamrup Metropolitan (district code 322)
3. Verifies pc11_tv_id matches Census Town/Village codes
4. Generates an inspection report

Usage:
    python inspect_aikosh_geometry.py [--input PATH] [--output PATH]

Expected input:
    - village_modified.shp (ESRI Shapefile) OR
    - village_modified.gpkg (GeoPackage)
    Located in: data/raw/aikosh_census_village_geometry_2011/
"""

import csv
import json
import os
import sys
from pathlib import Path
from datetime import datetime

# Try to import geopandas
try:
    import geopandas as gpd
    HAS_GEOPANDAS = True
except ImportError:
    HAS_GEOPANDAS = False

# Try to import openpyxl
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DATA_RAW = PROJECT_ROOT / "data" / "raw"
DATA_INSPECTION = PROJECT_ROOT / "data" / "inspection"
GEOMETRY_DIR = DATA_RAW / "aikosh_census_village_geometry_2011"
CENSUS_FILE = DATA_RAW / "kamrup_metropolitan_pca_tv_2011.xlsx"

# Kamrup Metropolitan codes
ASSAM_STATE_CODE = "18"
KAMRUP_METRO_DISTRICT_CODE = "322"  # Census PCA-TV district code
KAMRUP_METRO_DISTRICT_NAME = "Kamrup Metropolitan"


def load_census_data():
    """Load Census PCA-TV village codes for Kamrup Metropolitan."""
    print("Loading Census PCA-TV dataset...")

    if not HAS_OPENPYXL:
        print("  ERROR: openpyxl not installed. Run: pip install openpyxl")
        return {}

    if not CENSUS_FILE.exists():
        print(f"  ERROR: Census file not found: {CENSUS_FILE}")
        return {}

    wb = openpyxl.load_workbook(str(CENSUS_FILE), read_only=True, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    tv_col = headers.index('Town/Village') if 'Town/Village' in headers else None
    name_col = headers.index('Name') if 'Name' in headers else None
    level_col = headers.index('Level') if 'Level' in headers else None
    dist_col = headers.index('District') if 'District' in headers else None
    tru_col = headers.index('TRU') if 'TRU' in headers else None
    tot_p_col = headers.index('TOT_P') if 'TOT_P' in headers else None

    census_villages = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[tv_col] is not None:
            tv_code = str(int(row[tv_col])) if isinstance(row[tv_col], (int, float)) else str(row[tv_col])
            name = row[name_col] if name_col else ""
            level = row[level_col] if level_col else ""
            dist = row[dist_col] if dist_col else ""
            tru = row[tru_col] if tru_col else ""
            tot_p = row[tot_p_col] if tot_p_col else 0

            # Only include villages (not wards, sub-districts, etc.)
            if level == "VILLAGE" or level == "TOWN":
                census_villages[tv_code] = {
                    "name": name,
                    "level": level,
                    "district": dist,
                    "tru": tru,
                    "tot_p": tot_p,
                }

    wb.close()
    print(f"  Loaded {len(census_villages)} village/town records from Census")
    return census_villages


def find_geometry_file():
    """Find the village geometry file (shapefile or GeoPackage)."""
    print("\nLooking for geometry files...")

    if not GEOMETRY_DIR.exists():
        print(f"  Directory not found: {GEOMETRY_DIR}")
        return None, None

    # Check for GeoPackage first (single file)
    gpkg_files = list(GEOMETRY_DIR.glob("*.gpkg"))
    if gpkg_files:
        print(f"  Found GeoPackage: {gpkg_files[0].name}")
        return gpkg_files[0], "gpkg"

    # Check for shapefile components
    shp_files = list(GEOMETRY_DIR.glob("*.shp"))
    if shp_files:
        print(f"  Found Shapefile: {shp_files[0].name}")
        return shp_files[0], "shp"

    # List what's in the directory
    print(f"  Files in {GEOMETRY_DIR}:")
    for f in sorted(GEOMETRY_DIR.iterdir()):
        print(f"    - {f.name} ({f.stat().st_size:,} bytes)")

    return None, None


def inspect_geometry(geometry_file, file_type, census_villages):
    """Inspect the geometry file and verify matches with Census data."""
    print(f"\nInspecting geometry file: {geometry_file.name}")

    if not HAS_GEOPANDAS:
        print("  ERROR: geopandas not installed. Run: pip install geopandas")
        return None

    # Load geometry data
    print("  Loading geometry data...")
    if file_type == "gpkg":
        # List available layers
        layers = gpd.list_layers(str(geometry_file))
        print(f"  Available layers: {[l.name for l in layers]}")

        # Try to load the village layer
        gdf = gpd.read_file(str(geometry_file), layer="village_modified")
    else:
        gdf = gpd.read_file(str(geometry_file))

    print(f"  Total features: {len(gdf)}")
    print(f"  Columns: {list(gdf.columns)}")
    print(f"  CRS: {gdf.crs}")

    # Check for PC11 village ID columns
    possible_id_cols = [
        'pc11_tv_id', 'pc11_town_village_id', 'tv_id',
        'pc11_village_id', 'village_id', 'CEN_2011',
        'pc11_tv', 'pc11_s_id', 'pc11_d_id', 'pc11_sd_id'
    ]

    id_col = None
    for col in possible_id_cols:
        if col in gdf.columns:
            id_col = col
            print(f"  Found village ID column: {col}")
            break

    if id_col is None:
        print("  WARNING: No standard village ID column found")
        print("  Available columns:", list(gdf.columns))
        # Try to find any column that might contain village codes
        for col in gdf.columns:
            sample = gdf[col].dropna().head(5).tolist()
            if any(str(v).startswith('303') for v in sample):
                print(f"  Possible village code column: {col} (sample: {sample[:3]})")
                id_col = col
                break

    # Check for district ID column
    dist_col = None
    for col in ['pc11_d_id', 'pc11_district_id', 'district_id', 'DISTRICT']:
        if col in gdf.columns:
            dist_col = col
            print(f"  Found district ID column: {col}")
            break

    # Filter to Kamrup Metropolitan if district column exists
    kamrup_gdf = gdf
    if dist_col:
        kamrup_gdf = gdf[gdf[dist_col].astype(str) == KAMRUP_METRO_DISTRICT_CODE]
        print(f"  Features in Kamrup Metro (district {KAMRUP_METRO_DISTRICT_CODE}): {len(kamrup_gdf)}")

    # Verify matches with Census data
    if id_col:
        print(f"\n  Verifying village code matches...")
        geometry_ids = set(kamrup_gdf[id_col].astype(str).tolist())
        census_ids = set(census_villages.keys())

        matched = geometry_ids & census_ids
        only_geometry = geometry_ids - census_ids
        only_census = census_ids - geometry_ids

        print(f"  Matched: {len(matched)}")
        print(f"  Only in geometry: {len(only_geometry)}")
        print(f"  Only in Census: {len(only_census)}")

        if matched:
            print(f"\n  Sample matched villages:")
            for code in sorted(matched)[:5]:
                census_info = census_villages[code]
                print(f"    {code}: Census='{census_info['name']}' Geometry='{code}'")

        if only_geometry:
            print(f"\n  Village codes only in geometry (first 5):")
            for code in sorted(only_geometry)[:5]:
                print(f"    {code}")

        if only_census:
            print(f"\n  Village codes only in Census (first 5):")
            for code in sorted(only_census)[:5]:
                info = census_villages[code]
                print(f"    {code}: {info['name']} ({info['level']})")

        match_rate = len(matched) / len(census_ids) * 100 if census_ids else 0
        print(f"\n  Match rate: {match_rate:.1f}% ({len(matched)}/{len(census_ids)})")

        return {
            "total_features": len(gdf),
            "kamrup_features": len(kamrup_gdf),
            "id_column": id_col,
            "district_column": dist_col,
            "matched": len(matched),
            "only_geometry": len(only_geometry),
            "only_census": len(only_census),
            "match_rate": match_rate,
            "sample_matched": sorted(matched)[:10],
            "sample_only_geometry": sorted(only_geometry)[:10],
            "sample_only_census": sorted(only_census)[:10],
            "columns": list(gdf.columns),
            "crs": str(gdf.crs),
        }

    return None


def generate_report(inspection_results, census_villages):
    """Generate the inspection report."""
    print("\nGenerating inspection report...")

    report_path = DATA_INSPECTION / "aikosh_census_geometry_inspection.md"

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    report = f"""# Aikosh/SHRUG PC11 Village Geometry Inspection Report

**Generated**: {timestamp}
**Dataset**: 2011 Population Census Village-Level Geometries (PC11 Village Polygons)
**Source**: AIKOSH/SHRUG (Development Data Lab)
**Target**: Kamrup Metropolitan, Assam (District code {KAMRUP_METRO_DISTRICT_CODE})

## Summary

| Metric | Value |
|--------|-------|
| Total features in geometry file | {inspection_results.get('total_features', 'N/A')} |
| Features in Kamrup Metropolitan | {inspection_results.get('kamrup_features', 'N/A')} |
| Village ID column | `{inspection_results.get('id_column', 'N/A')}` |
| District ID column | `{inspection_results.get('district_column', 'N/A')}` |
| Matched village codes | {inspection_results.get('matched', 'N/A')} |
| Only in geometry | {inspection_results.get('only_geometry', 'N/A')} |
| Only in Census | {inspection_results.get('only_census', 'N/A')} |
| Match rate | {inspection_results.get('match_rate', 'N/A'):.1f}% |
| CRS | {inspection_results.get('crs', 'N/A')} |

## Verification Result

"""

    match_rate = inspection_results.get('match_rate', 0)
    if match_rate and match_rate >= 90:
        report += "**PASS**: Village geometry can be joined to Census PCA-TV data via village codes.\n"
    elif match_rate and match_rate >= 70:
        report += "**PARTIAL**: Some village codes match. Manual verification recommended.\n"
    else:
        report += "**FAIL**: Village codes do not match. Different coding systems may be in use.\n"

    report += f"""
## Census Data Reference

- **File**: `data/raw/kamrup_metropolitan_pca_tv_2011.xlsx`
- **District code**: {KAMRUP_METRO_DISTRICT_CODE} (Kamrup Metropolitan)
- **Total village/town records**: {len(census_villages)}
- **Village codes are 6-digit numbers** (e.g., 303398, 303399, etc.)

## Geometry Data Reference

- **Dataset**: 2011 Population Census Village-Level Geometries
- **Source**: AIKOSH (https://aikosh.indiaai.gov.in/home/datasets/details/2011_population_census_village_level_geometries.html)
- **Alternative**: DDL (https://www.devdatalab.org/shrug_download)
- **License**: CC BY-NC-SA 4.0
- **Expected columns**: pc11_tv_id (village code), pc11_d_id (district code), geometry

## Matched Village Codes

"""

    if inspection_results.get('sample_matched'):
        report += "| Code | Census Name |\n|------|-------------|\n"
        for code in inspection_results['sample_matched']:
            name = census_villages.get(code, {}).get('name', 'Unknown')
            report += f"| {code} | {name} |\n"

    report += f"""

## Files in Directory

```
data/raw/aikosh_census_village_geometry_2011/
├── village_modified.shp      (geometry)
├── village_modified.shx      (index)
├── village_modified.dbf      (attributes)
├── village_modified.prj      (projection)
├── village_modified.cpg      (encoding)
├── README.md                 (metadata)
└── open_poly.bib             (citation)
```

## Next Steps

1. Filter geometry to Kamrup Metropolitan using district code {KAMRUP_METRO_DISTRICT_CODE}
2. Join with Census PCA-TV data using village codes
3. Export filtered GeoJSON/Shapefile for frontend use
4. Update `docs/data-requirements/kamrup_metro_village_geometry.md`

## Columns in Geometry File

{json.dumps(inspection_results.get('columns', []), indent=2)}
"""

    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report)

    print(f"  Report saved to: {report_path}")
    return report_path


def main():
    print("=" * 70)
    print("Aikosh/SHRUG PC11 Village Geometry Inspection")
    print("Target: Kamrup Metropolitan, Assam")
    print("=" * 70)

    # Load Census data
    census_villages = load_census_data()
    if not census_villages:
        print("\nERROR: Could not load Census data")
        sys.exit(1)

    # Find geometry file
    geometry_file, file_type = find_geometry_file()
    if not geometry_file:
        print("\n" + "=" * 70)
        print("GEOMETRY FILE NOT FOUND")
        print("=" * 70)
        print(f"""
Please download the village geometry data to:
    {GEOMETRY_DIR}/

Download sources:
1. AIKOSH (requires registration):
   https://aikosh.indiaai.gov.in/home/datasets/details/2011_population_census_village_level_geometries.html

2. DDL (requires JavaScript interaction):
   https://www.devdatalab.org/shrug_download

Files to download:
    village_modified.shp
    village_modified.shx
    village_modified.dbf
    village_modified.prj
    village_modified.cpg

After downloading, re-run this script.
""")
        sys.exit(1)

    # Inspect geometry
    inspection_results = inspect_geometry(geometry_file, file_type, census_villages)

    if inspection_results:
        # Generate report
        report_path = generate_report(inspection_results, census_villages)

        print("\n" + "=" * 70)
        print("INSPECTION COMPLETE")
        print("=" * 70)
        print(f"  Report: {report_path}")
        print(f"  Match rate: {inspection_results['match_rate']:.1f}%")
    else:
        print("\nERROR: Could not inspect geometry file")
        sys.exit(1)


if __name__ == "__main__":
    main()
