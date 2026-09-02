"""
Census code verification - compare AIKOSH geometry with Census PCA-TV data.
Read-only comparison for Kamrup Metropolitan (District 322).
"""
import json
import sys
from pathlib import Path

try:
    import geopandas as gpd
except ImportError:
    print("ERROR: geopandas not installed")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed")
    sys.exit(1)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
SHAPEFILE = PROJECT_ROOT / "data/raw/aikosh_census_village_geometry_2011/extracted/shrug-pc11-village-poly-shp/village_modified.shp"
CENSUS_FILE = PROJECT_ROOT / "data/raw/kamrup_metropolitan_pca_tv_2011.xlsx"
OUTPUT_JSON = PROJECT_ROOT / "data/inspection/aikosh_census_verification.json"

SEP = "=" * 70
LINE = "-" * 70

print(SEP)
print("CENSUS CODE VERIFICATION - Kamrup Metropolitan (District 322)")
print(SEP)

# ── Load Census data ──
print()
print("Loading Census PCA-TV data...")
wb = openpyxl.load_workbook(str(CENSUS_FILE), read_only=True, data_only=True)
ws = wb.active

headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print("  Census columns: " + str(headers[:15]))

# Find column indices
tv_col = headers.index('Town/Village') if 'Town/Village' in headers else None
name_col = headers.index('Name') if 'Name' in headers else None
level_col = headers.index('Level') if 'Level' in headers else None
dist_col = headers.index('District') if 'District' in headers else None
subdist_col = headers.index('Subdistt') if 'Subdistt' in headers else None

# Read Census data
census_all = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[tv_col] is not None:
        tv_code = str(int(row[tv_col])) if isinstance(row[tv_col], (int, float)) else str(row[tv_col])
        name = row[name_col] if name_col else ""
        level = row[level_col] if level_col else ""
        dist = str(int(row[dist_col])) if isinstance(row[dist_col], (int, float)) else str(row[dist_col]) if row[dist_col] else ""
        subdist = str(int(row[subdist_col])) if isinstance(row[subdist_col], (int, float)) else str(row[subdist_col]) if row[subdist_col] else ""
        census_all.append({
            "tv_code": tv_code,
            "name": name,
            "level": level,
            "district": dist,
            "subdistt": subdist,
        })

wb.close()

# Filter to District 322 (Kamrup Metro)
census_322 = [r for r in census_all if r["district"] == "322"]
print(f"  Total Census records: {len(census_all)}")
print(f"  Census records District 322: {len(census_322)}")

# Show Census District 322 levels
levels_322 = {}
for r in census_322:
    lvl = r["level"]
    levels_322[lvl] = levels_322.get(lvl, 0) + 1
print("  Census District 322 by Level: " + str(levels_322))

# Filter to villages only (VILLAGE and TOWN)
census_villages_322 = [r for r in census_322 if r["level"] in ("VILLAGE", "TOWN")]
print(f"  Census villages/towns District 322: {len(census_villages_322)}")

# ── Load AIKOSH data ──
print()
print("Loading AIKOSH shapefile...")
gdf = gpd.read_file(str(SHAPEFILE))
print(f"  Total AIKOSH features: {len(gdf):,}")

# Filter to District 322
kamrup_mask = gdf['pc11_d_id'].astype(str) == '322'
kamrup_gdf = gdf[kamrup_mask].copy()
print(f"  AIKOSH features District 322: {len(kamrup_gdf):,}")

# ── Compare codes ──
print()
print(LINE)
print("PC11_TV_ID MATCHING ANALYSIS")
print(LINE)

# Census village codes (District 322)
census_tv_codes = set(r["tv_code"] for r in census_villages_322)
# AIKOSH village codes (District 322)
aikosh_tv_codes = set(kamrup_gdf['pc11_tv_id'].astype(str).tolist())

matched_codes = census_tv_codes & aikosh_tv_codes
only_census = census_tv_codes - aikosh_tv_codes
only_aikosh = aikosh_tv_codes - census_tv_codes

print(f"  Census village/town codes (District 322): {len(census_tv_codes)}")
print(f"  AIKOSH pc11_tv_id codes (District 322): {len(aikosh_tv_codes)}")
print(f"  EXACT MATCHES: {len(matched_codes)}")
print(f"  Only in Census (missing geometry): {len(only_census)}")
print(f"  Only in AIKOSH (extra geometry): {len(only_aikosh)}")

# Match rate
match_rate = len(matched_codes) / len(census_tv_codes) * 100 if census_tv_codes else 0
print(f"  MATCH RATE: {match_rate:.1f}%")

# ── Duplicate check ──
print()
print(LINE)
print("DUPLICATE CHECK")
print(LINE)
tv_id_counts = kamrup_gdf['pc11_tv_id'].astype(str).value_counts()
duplicates = tv_id_counts[tv_id_counts > 1]
print(f"  Duplicate pc11_tv_id values: {len(duplicates)}")
if len(duplicates) > 0:
    for tv_id, count in duplicates.head(10).items():
        rows = kamrup_gdf[kamrup_gdf['pc11_tv_id'].astype(str) == tv_id]
        names = rows['tv_name'].tolist()
        print(f"    {tv_id}: {count} features, names={names}")

# ── Name comparison ──
print()
print(LINE)
print("NAME COMPARISON (matched codes)")
print(LINE)

# Build lookup dictionaries
census_by_code = {r["tv_code"]: r["name"] for r in census_villages_322}
aikosh_by_code = {}
for _, row in kamrup_gdf.iterrows():
    code = str(row['pc11_tv_id'])
    aikosh_by_code[code] = row['tv_name']

name_matches = 0
name_mismatches = 0
name_mismatch_details = []

for code in sorted(matched_codes):
    census_name = str(census_by_code.get(code, "")).strip()
    aikosh_name = str(aikosh_by_code.get(code, "")).strip()
    if census_name.lower() == aikosh_name.lower():
        name_matches += 1
    else:
        name_mismatches += 1
        name_mismatch_details.append({
            "code": code,
            "census_name": census_name,
            "aikosh_name": aikosh_name,
        })

print(f"  Names match: {name_matches}")
print(f"  Names differ: {name_mismatches}")
if name_mismatch_details:
    print()
    print("  Name mismatches (first 10):")
    for m in name_mismatch_details[:10]:
        print(f"    {m['code']}: Census='{m['census_name']}' vs AIKOSH='{m['aikosh_name']}'")

# ── Geometry validation ──
print()
print(LINE)
print("GEOMETRY VALIDATION (District 322)")
print(LINE)
valid_geom = kamrup_gdf.geometry.is_valid.sum()
invalid_geom = (~kamrup_gdf.geometry.is_valid).sum()
missing_geom = kamrup_gdf.geometry.isna().sum()
print(f"  Valid polygon geometry: {valid_geom:,}")
print(f"  Invalid geometry: {invalid_geom:,}")
print(f"  Missing geometry: {missing_geom:,}")

# ── Subdistt comparison ──
print()
print(LINE)
print("SUBDISTT COMPARISON")
print(LINE)

# Census subdistt codes for District 322
census_subdistts = set(r["subdistt"] for r in census_villages_322 if r["subdistt"])
aikosh_subdistts = set(kamrup_gdf['pc11_sd_id'].astype(str).unique().tolist())

print(f"  Census Subdistt codes (District 322): {sorted(census_subdistts)}")
print(f"  AIKOSH pc11_sd_id codes (District 322): {sorted(aikosh_subdistts)}")
subdistt_match = census_subdistts == aikosh_subdistts
print(f"  Subdistt codes match: {'YES' if subdistt_match else 'NO'}")

# ── Sample matched villages ──
print()
print(LINE)
print("SAMPLE MATCHED VILLAGES (first 20)")
print(LINE)
print(f"  {'Code':>8s}  {'Census Name':30s}  {'AIKOSH Name':30s}  {'Match':>5s}")
print(f"  {'----':>8s}  {'-'*30:30s}  {'-'*30:30s}  {'-----':>5s}")
for code in sorted(matched_codes)[:20]:
    census_name = str(census_by_code.get(code, ""))
    aikosh_name = str(aikosh_by_code.get(code, ""))
    name_ok = "YES" if census_name.lower() == aikosh_name.lower() else "NO"
    print(f"  {code:>8s}  {census_name:30s}  {aikosh_name:30s}  {name_ok:>5s}")

# ── Unmatched Census villages ──
print()
print(LINE)
print("UNMATCHED CENSUS VILLAGES (missing geometry)")
print(LINE)
if only_census:
    for code in sorted(only_census):
        name = census_by_code.get(code, "")
        level = next((r["level"] for r in census_villages_322 if r["tv_code"] == code), "")
        print(f"  {code}: {name} ({level})")
else:
    print("  None - all Census villages have matching AIKOSH geometry")

# ── Save results ──
results = {
    "census_total_records": len(census_all),
    "census_district_322_total": len(census_322),
    "census_district_322_villages": len(census_villages_322),
    "census_district_322_levels": levels_322,
    "aikosh_total_features": len(gdf),
    "aikosh_district_322_features": len(kamrup_gdf),
    "pc11_tv_id_matches": len(matched_codes),
    "only_in_census": len(only_census),
    "only_in_aikosh": len(only_aikosh),
    "match_rate_pct": round(match_rate, 1),
    "duplicate_pc11_tv_id": len(duplicates),
    "name_matches": name_matches,
    "name_mismatches": name_mismatches,
    "name_mismatch_details": name_mismatch_details,
    "valid_geometry": int(valid_geom),
    "invalid_geometry": int(invalid_geom),
    "missing_geometry": int(missing_geom),
    "census_subdistt_codes": sorted(list(census_subdistts)),
    "aikosh_subdistt_codes": sorted(list(aikosh_subdistts)),
    "unmatched_census_codes": sorted(list(only_census)),
}

with open(OUTPUT_JSON, "w") as f:
    json.dump(results, f, indent=2, default=str)

print()
print(SEP)
print("CENSUS VERIFICATION COMPLETE")
print("Results saved to: " + str(OUTPUT_JSON))
print(SEP)
